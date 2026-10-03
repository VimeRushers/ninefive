"""Parse uploaded pricelists (CSV, XLSX, PDF) into catalogue items.

Header names are matched in Romanian, English and Russian. Files without a
recognizable header fall back to the column order [name, description, price].
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path


@dataclass
class ParsedItem:
    name: str
    description: str
    price_amount: float | None
    price_currency: str


NAME_HEADERS = {
    "name",
    "nume",
    "denumire",
    "denumirea",
    "produs",
    "articol",
    "item",
    "наименование",
}
DESCRIPTION_HEADERS = {
    "description",
    "descriere",
    "descrierea",
    "specificatie",
    "specificație",
    "detalii",
    "описание",
}
PRICE_HEADERS = {
    "price",
    "pret",
    "preț",
    "cost",
    "valoare",
    "valoarea",
    "unit price",
    "unit_price",
    "цена",
}


def parse_pricelist(path: str | Path) -> list[ParsedItem]:
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix == ".csv":
        rows = _read_csv(path)
    elif suffix in {".xlsx", ".xls"}:
        rows = _read_excel(path)
    elif suffix == ".pdf":
        rows = _read_pdf(path)
    else:
        return []

    return _rows_to_items(rows)


def _read_csv(path: Path) -> list[list[str]]:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [[cell for cell in row] for row in csv.reader(io.StringIO(text), dialect)]


def _read_excel(path: Path) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    if sheet is None:
        workbook.close()
        return []
    rows = [
        ["" if cell is None else str(cell) for cell in row]
        for row in sheet.iter_rows(values_only=True)
    ]
    workbook.close()
    return rows


def _read_pdf(path: Path) -> list[list[str]]:
    try:
        import pdfplumber
    except ImportError:
        return []

    rows: list[list[str]] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables() or []:
                rows.extend(
                    ["" if cell is None else str(cell) for cell in row] for row in table
                )
    return rows


def _normalize(value: object) -> str:
    return str(value).strip().lower() if value is not None else ""


def _find_header(rows: list[list[str]]) -> tuple[int, dict[str, int]] | None:
    for index, row in enumerate(rows[:5]):
        columns: dict[str, int] = {}
        for position, cell in enumerate(_normalize(c) for c in row):
            if cell in NAME_HEADERS and "name" not in columns:
                columns["name"] = position
            elif cell in DESCRIPTION_HEADERS and "description" not in columns:
                columns["description"] = position
            elif cell in PRICE_HEADERS and "price" not in columns:
                columns["price"] = position
        if "name" in columns and "price" in columns:
            return index, columns
    return None


def _rows_to_items(rows: list[list[str]]) -> list[ParsedItem]:
    if not rows:
        return []

    header = _find_header(rows)
    if header is None:
        start = 0
        columns = {"name": 0, "description": 1, "price": 2}
    else:
        start = header[0] + 1
        columns = header[1]

    items: list[ParsedItem] = []
    for row in rows[start:]:
        if not row:
            continue

        def cell(key: str) -> str:
            position = columns.get(key)
            if position is None or position >= len(row):
                return ""
            return str(row[position]).strip()

        name = cell("name")
        if not name:
            continue

        raw_price = cell("price")
        items.append(
            ParsedItem(
                name=name,
                description=cell("description"),
                price_amount=_parse_price(raw_price),
                price_currency=_parse_currency(raw_price),
            )
        )
    return items


def _parse_price(raw: str) -> float | None:
    text = re.sub(r"[^\d,.\-]", "", str(raw))
    if not text:
        return None

    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")

    try:
        return float(text)
    except ValueError:
        return None


def _parse_currency(raw: str) -> str:
    upper = str(raw).upper()
    for code in ("MDL", "EUR", "USD", "RON"):
        if code in upper:
            return code
    return "MDL"
